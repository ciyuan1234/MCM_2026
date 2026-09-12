"""提交前交付自检：逐项核对四个 result*.xlsx、manifest、图件与文档。

只读检查，不重算、不修改任何产物。用法：

    python analysis/delivery_check.py            # 打印 Markdown 表格
    python analysis/delivery_check.py --strict   # 有 FAIL 时以非零状态退出
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
import sys
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from PIL import Image

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

OUTPUTS = PROJECT_ROOT / "outputs"
TEMPLATES = PROJECT_ROOT / "附件" / "附件3"

EXPECTED = {
    1: {
        "sheets": ("温度", "水分浓度"),
        "step_s": 1.0,
        "last_s": 1800.0,
        "rows": 1802,
        "columns": 22,
        "last_header": None,
    },
    2: {
        "sheets": ("温度", "水分浓度"),
        "step_s": 1.0,
        "last_s": 10800.0,
        "rows": 10802,
        "columns": 22,
        "last_header": None,
    },
    3: {
        "sheets": ("Sheet1",),
        "step_s": 60.0,
        "last_s": 206160.0,
        "rows": 3438,
        "columns": 22,
        "last_header": None,
    },
    4: {
        "sheets": ("Sheet1",),
        "step_s": 60.0,
        "last_s": 183060.0,
        "rows": 3053,
        "columns": 22,
        "last_header": "药材表面",
    },
}
MAX_AIR_TEMPERATURE_C = 50.165

FIGURES = (
    "fig1_q1_temperature_profiles",
    "fig2_q1_moisture_profiles",
    "fig3_q1_center_surface_temperature",
    "fig4_q1_center_surface_moisture",
    "fig5_q1_summary_2x2",
    "fig1_q2_temperature_profiles",
    "fig2_q2_moisture_profiles",
    "fig3_q2_center_surface_temperature",
    "fig4_q2_center_surface_moisture",
    "fig5_q2_summary_2x2",
    "fig1_q3_temperature_profiles",
    "fig2_q3_moisture_profiles",
    "fig3_q3_center_surface_temperature",
    "fig4_q3_center_surface_moisture",
    "fig5_q3_summary_2x2",
    "fig6_q3_convergence",
    "fig1_q4_radius_history",
    "fig2_q4_moisture_profiles",
    "fig3_q4_center_surface_moisture",
    "fig4_q4_center_surface_temperature",
    "fig5_q4_summary_2x2",
    "fig6_q4_q3_comparison",
)

DOCUMENTS = (
    "AGENTS.md",
    "docs/project_plan.md",
    "docs/overview.md",
    "docs/decision_log.md",
    "docs/data_dictionary.md",
    "docs/data_processing.md",
    "docs/gap_register.md",
    "docs/verification_matrix.md",
    "docs/requirements_traceability.md",
    "docs/notation.md",
    "docs/paper_outline.md",
    "docs/paper_figures.md",
    "docs/paper_tables.md",
    "docs/visualization_spec.md",
    "docs/problem_delivery_standard.md",
) + tuple(
    f"docs/problem_{problem}_{suffix}.md"
    for problem in (1, 2, 3, 4)
    for suffix in ("model", "decision_log", "paper_materials")
)

OVERVIEW_NUMBERS = (
    "57.2667",
    "50.8500",
    "6.4167",
    "1.5103",
    "1.7662",
    "0.0539",
    "0.0058",
    "3.754e-4",
    "5.26e-5",
)

# 正文草稿使用排版式科学计数法（如 3.754×10⁻⁴），检查时先归一化再比较数值。
SUPERSCRIPT_TABLE = str.maketrans("⁰¹²³⁴⁵⁶⁷⁸⁹⁻", "0123456789-")
DRAFT_NUMBERS = (57.2667, 50.85, 6.4167, 1.5103, 1.7662, 0.0539, 0.0058, 3.754e-4, 5.26e-5)


def _normalized_numbers(text: str) -> list[float]:
    normalized = text.translate(SUPERSCRIPT_TABLE).replace("×10", "e")
    values = []
    for token in re.findall(r"\d+\.?\d*(?:e-?\d+)?", normalized):
        try:
            values.append(float(token))
        except ValueError:  # pragma: no cover
            continue
    return values


def _record(results: list[dict], name: str, ok: bool, detail: str) -> None:
    results.append({"name": name, "status": "PASS" if ok else "FAIL", "detail": detail})


def _data_matrix(sheet) -> tuple[np.ndarray, np.ndarray]:
    rows = list(sheet.iter_rows(values_only=True))
    header = rows[0]
    body = rows[1:]
    times = np.array([row[0] for row in body], dtype=float)
    values = np.array([row[1:] for row in body], dtype=float)
    return times, values, header


def _check_workbook(problem: int, results: list[dict]) -> None:
    expected = EXPECTED[problem]
    path = OUTPUTS / f"result{problem}.xlsx"
    template_path = TEMPLATES / f"result{problem}.xlsx"
    if not path.exists():
        _record(results, f"问题 {problem} 结果文件存在", False, f"缺少 {path.name}")
        return
    workbook = load_workbook(path)
    template = load_workbook(template_path, read_only=True)
    try:
        sheet_names = tuple(workbook.sheetnames)
        _record(
            results,
            f"问题 {problem} sheet 结构",
            sheet_names == expected["sheets"],
            f"{sheet_names} vs 期望 {expected['sheets']}",
        )
        template_names = tuple(template.sheetnames)
        _record(
            results,
            f"问题 {problem} 与模板 sheet 一致",
            sheet_names == template_names,
            f"{sheet_names} vs 模板 {template_names}",
        )
        for name in expected["sheets"]:
            sheet = workbook[name]
            times, values, header = _data_matrix(sheet)
            label = f"问题 {problem} [{name}]"
            _record(
                results,
                f"{label} 行列数与模板一致",
                sheet.max_row == expected["rows"] and sheet.max_column == expected["columns"],
                f"{sheet.max_row} 行 × {sheet.max_column} 列，期望 "
                f"{expected['rows']} × {expected['columns']}",
            )
            _record(
                results,
                f"{label} 表头",
                header[0] == "时间\\到药材中心的距离"
                and (
                    expected["last_header"] is None
                    or header[-1] == expected["last_header"]
                ),
                f"A1={header[0]!r}，末列={header[-1]!r}",
            )
            steps = np.diff(times)
            _record(
                results,
                f"{label} 时间列",
                times[0] == 0.0
                and abs(times[-1] - expected["last_s"]) < 1e-9
                and np.allclose(steps, expected["step_s"], atol=1e-9),
                f"0 → {times[-1]:.0f} s，步长 {steps[0]:.0f} s",
            )
            _record(
                results,
                f"{label} 数值有限且四位小数",
                bool(np.all(np.isfinite(values)))
                and all(
                    cell.number_format == "0.0000"
                    for row in sheet.iter_rows(
                        min_row=2, min_col=2, max_col=sheet.max_column
                    )
                    for cell in row
                ),
                "无缺失/非有限值；数据区格式 0.0000",
            )
            moisture_like = values
            _record(
                results,
                f"{label} 数值范围合理",
                bool(np.all(moisture_like >= 0.0)),
                f"最小值 {float(moisture_like.min()):.4f}",
            )
    finally:
        workbook.close()
        template.close()


def _check_moisture_semantics(problem: int, results: list[dict]) -> None:
    """问题 3、4 的单表为含水率：单调不增、圆心最大、末值 0.1500。"""
    path = OUTPUTS / f"result{problem}.xlsx"
    if not path.exists():
        return
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[EXPECTED[problem]["sheets"][0]]
        times, values, _ = _data_matrix(sheet)
    finally:
        workbook.close()
    tolerance = 1e-9
    _record(
        results,
        f"问题 {problem} 含水率随时间单调不增",
        bool(np.all(np.diff(values, axis=0) <= tolerance)),
        "逐列检查",
    )
    _record(
        results,
        f"问题 {problem} 圆心为空间最大含水率",
        bool(np.allclose(values.max(axis=1), values[:, 0], atol=1e-4)),
        "逐行检查",
    )
    _record(
        results,
        f"问题 {problem} 末时刻圆心含水率命中阈值",
        abs(float(values[-1, 0]) - 0.1500) < 5e-5,
        f"{float(values[-1, 0]):.4f}",
    )
    _record(
        results,
        f"问题 {problem} 末时刻时间与判据一致",
        abs(float(times[-1]) - EXPECTED[problem]["last_s"]) < 1e-9,
        f"{float(times[-1]):.0f} s",
    )


def _check_temperature_bound(results: list[dict]) -> None:
    maximum = -np.inf
    for problem in (1, 2):
        workbook = load_workbook(OUTPUTS / f"result{problem}.xlsx", read_only=True, data_only=True)
        try:
            sheet = workbook["温度"]
            _, values, _ = _data_matrix(sheet)
            maximum = max(maximum, float(values.max()))
        finally:
            workbook.close()
    _record(
        results,
        "温度不超过空气温度上界",
        maximum <= MAX_AIR_TEMPERATURE_C + 0.01,
        f"最高 {maximum:.4f} °C（空气末值 {MAX_AIR_TEMPERATURE_C} °C）",
    )


def _check_attachments_read_only(results: list[dict]) -> None:
    try:
        completed = subprocess.run(
            ["git", "status", "--porcelain", "--", "附件"],
            cwd=PROJECT_ROOT,
            capture_output=True,
            text=True,
            check=True,
        )
        _record(
            results,
            "附件目录未被修改",
            completed.stdout.strip() == "",
            "git status 干净" if completed.stdout.strip() == "" else completed.stdout.strip(),
        )
    except (subprocess.CalledProcessError, FileNotFoundError) as exc:  # pragma: no cover
        _record(results, "附件目录未被修改", False, str(exc))
    digests = []
    for name in ("附件1.xlsx", "附件2.xlsx"):
        path = PROJECT_ROOT / "附件" / name
        digest = hashlib.sha256(path.read_bytes()).hexdigest()[:12]
        digests.append(f"{name}:{digest}")
    for index in (1, 2, 3, 4):
        path = TEMPLATES / f"result{index}.xlsx"
        digests.append(f"result{index}.xlsx:{hashlib.sha256(path.read_bytes()).hexdigest()[:12]}")
    _record(results, "附件与模板哈希已记录", True, "；".join(digests))


def _check_manifests(results: list[dict]) -> None:
    required = ("problem", "timestamp", "command", "inputs", "outputs", "parameters", "code_hashes", "environment")
    for problem in (1, 2, 3, 4):
        path = OUTPUTS / f"run_manifest_q{problem}.json"
        if not path.exists():
            _record(results, f"问题 {problem} manifest 存在", False, f"缺少 {path.name}")
            continue
        with path.open(encoding="utf-8") as file:
            manifest = json.load(file)
        missing = [key for key in required if key not in manifest]
        _record(
            results,
            f"问题 {problem} manifest 字段完整",
            not missing,
            "字段齐全" if not missing else f"缺少 {missing}",
        )
    latest = OUTPUTS / "run_manifest.json"
    _record(results, "最新 run_manifest.json 存在", latest.exists(), latest.name)


def _check_headline_numbers(results: list[dict]) -> None:
    with (OUTPUTS / "run_manifest_q3.json").open(encoding="utf-8") as file:
        q3 = json.load(file)["parameters"]
    with (OUTPUTS / "run_manifest_q4.json").open(encoding="utf-8") as file:
        q4 = json.load(file)["parameters"]
    with (OUTPUTS / "q3_convergence.json").open(encoding="utf-8") as file:
        q3c = json.load(file)["comparisons"]
    with (OUTPUTS / "q4_convergence.json").open(encoding="utf-8") as file:
        q4c = json.load(file)["comparisons"]

    workbook1 = load_workbook(OUTPUTS / "result1.xlsx", read_only=True, data_only=True)
    try:
        _, values1, _ = _data_matrix(workbook1["水分浓度"])
    finally:
        workbook1.close()
    workbook2 = load_workbook(OUTPUTS / "result2.xlsx", read_only=True, data_only=True)
    try:
        _, values2, _ = _data_matrix(workbook2["水分浓度"])
    finally:
        workbook2.close()

    checks = {
        "问题 3 烘干时间": (round(float(q3["t_f_hours"]), 4), 57.2667),
        "问题 4 烘干时间": (round(float(q4["t_f_hours"]), 4), 50.85),
        "问题 4 相对问题 3 的差值": (
            round(float(q3["t_f_hours"]) - float(q4["t_f_hours"]), 4),
            6.4167,
        ),
        "问题 1 表面含水率 @1800 s": (round(float(values1[-1, -1]), 4), 1.5103),
        "问题 2 圆心含水率 @3 h": (round(float(values2[-1, 0]), 4), 1.7662),
        "问题 3 空间收敛": (round(float(q3c["spatial_t_f_difference_hours"]), 4), 0.0539),
        "问题 4 空间收敛": (round(float(q4c["spatial_t_f_difference_hours"]), 4), 0.0058),
    }
    for name, (actual, documented) in checks.items():
        _record(results, name, abs(actual - documented) < 5e-5, f"产物 {actual} / 文档 {documented}")

    _record(
        results,
        "问题 3 场差量级",
        abs(float(q3c["spatial_max_moisture_difference"]) - 3.754e-4) < 1e-6,
        f"{float(q3c['spatial_max_moisture_difference']):.3e}",
    )
    _record(
        results,
        "问题 4 场差量级",
        abs(float(q4c["spatial_max_moisture_difference"]) - 5.258e-5) < 1e-6,
        f"{float(q4c['spatial_max_moisture_difference']):.3e}",
    )

    overview = (PROJECT_ROOT / "docs" / "overview.md").read_text(encoding="utf-8")
    missing = [token for token in OVERVIEW_NUMBERS if token not in overview]
    _record(
        results,
        "总览包含全部头条数字",
        not missing,
        "全部命中" if not missing else f"缺少 {missing}",
    )


def _check_figures(results: list[dict]) -> None:
    missing = []
    empty = []
    for stem in FIGURES:
        for extension in ("png", "pdf", "svg"):
            path = OUTPUTS / "figures" / f"{stem}.{extension}"
            if not path.exists():
                missing.append(path.name)
            elif path.stat().st_size == 0:
                empty.append(path.name)
    _record(
        results,
        "图件齐全（22 张 × 3 格式）",
        not missing and not empty,
        "全部存在且非空" if not missing and not empty else f"缺少 {missing}；空文件 {empty}",
    )


def _check_documents(results: list[dict]) -> None:
    missing = [name for name in DOCUMENTS if not (PROJECT_ROOT / name).exists()]
    _record(
        results,
        "交付文档齐全",
        not missing,
        f"共 {len(DOCUMENTS)} 份" if not missing else f"缺少 {missing}",
    )
    experiments = sorted(path.parent.name for path in (PROJECT_ROOT / "docs" / "experiments").glob("*/brief.md"))
    _record(
        results,
        "实验任务卡齐全",
        len(experiments) >= 6,
        "、".join(experiments),
    )


def _check_paper_draft(results: list[dict]) -> None:
    """正文草稿必须存在，且头条数字与总览一致，防止论文与结果脱节。"""
    path = PROJECT_ROOT / "docs" / "paper_draft.md"
    if not path.exists():
        _record(results, "论文正文草稿存在", False, "缺少 docs/paper_draft.md")
        return
    text = path.read_text(encoding="utf-8")
    _record(results, "论文正文草稿存在", True, f"{len(text.splitlines())} 行")
    values = _normalized_numbers(text)
    missing = [
        expected
        for expected in DRAFT_NUMBERS
        if not any(abs(value - expected) <= max(abs(expected) * 1e-6, 1e-12) for value in values)
    ]
    _record(
        results,
        "正文草稿包含全部头条数字",
        not missing,
        "全部命中" if not missing else f"缺少 {missing}",
    )
    required_sections = (
        "摘要",
        "问题重述与总体分析",
        "数据预处理",
        "问题一：",
        "问题二：",
        "问题三：",
        "问题四：",
        "模型检验与敏感性汇总",
        "建模取舍与结论边界",
        "结论",
        "附录",
    )
    absent = [section for section in required_sections if section not in text]
    _record(
        results,
        "正文结构完整",
        not absent,
        "9 节 + 附录齐全" if not absent else f"缺少 {absent}",
    )
    steps = ("问题分析", "模型建立", "模型求解", "结果与分析")
    missing_steps = [
        f"{problem}.{index + 1} {step}"
        for problem in (3, 4, 5, 6)
        for index, step in enumerate(steps)
        if f"### {problem}.{index + 1} {step}" not in text
    ]
    _record(
        results,
        "四问均按“问题分析→模型建立→模型求解→结果”组织",
        not missing_steps,
        "四问结构一致" if not missing_steps else f"缺少 {missing_steps}",
    )
    image_paths = re.findall(r"!\[[^\]]*\]\(([^)]+)\)", text)
    missing_images = [
        target
        for target in image_paths
        if not (path.parent / target).resolve().exists()
    ]
    _record(
        results,
        "正文图件链接有效",
        bool(image_paths) and not missing_images,
        f"共 {len(image_paths)} 张，全部存在"
        if image_paths and not missing_images
        else f"失效链接 {missing_images}",
    )
    canonical = [target for target in image_paths if target.startswith("../outputs/figures/")]
    _record(
        results,
        "图件链接指向规范目录",
        len(canonical) == len(image_paths),
        f"{len(canonical)}/{len(image_paths)} 指向 ../outputs/figures/",
    )
    stems = [target.rsplit("/", 1)[-1].replace(".png", "") for target in image_paths]
    expected_stems = set(FIGURES)
    duplicated = sorted({stem for stem in stems if stems.count(stem) > 1})
    _record(
        results,
        "22 张图在正文与附录中各出现一次",
        not duplicated and set(stems) == expected_stems,
        f"共 {len(stems)} 张，重复 {duplicated}，缺失 {sorted(expected_stems - set(stems))}",
    )
    body_part, _, appendix_part = text.partition("## 附录 D")
    body_count = len(re.findall(r"!\[[^\]]*\]\(", body_part))
    appendix_count = len(re.findall(r"!\[[^\]]*\]\(", appendix_part))
    _record(
        results,
        "正文保留结论图、其余移入附录",
        body_count == 9 and appendix_count == 13,
        f"正文 {body_count} 张，附录 {appendix_count} 张",
    )


def _check_grayscale_figures(results: list[dict]) -> None:
    """黑白打印要求：所有图件必须是灰阶。"""
    worst = []
    for stem in FIGURES:
        path = OUTPUTS / "figures" / f"{stem}.png"
        if not path.exists():
            continue
        with Image.open(path) as image:
            array = np.asarray(image.convert("RGB"), dtype=float) / 255.0
        maximum = array.max(axis=2)
        minimum = array.min(axis=2)
        saturation = np.where(maximum > 0.0, (maximum - minimum) / np.maximum(maximum, 1e-9), 0.0)
        worst.append((float((saturation > 0.15).mean()), stem))
    worst.sort(reverse=True)
    offenders = [f"{stem}({fraction:.2%})" for fraction, stem in worst if fraction > 0.01]
    _record(
        results,
        "全部图件为黑白（灰阶）",
        not offenders,
        f"最高彩色像素占比 {worst[0][0]:.2%}（{worst[0][1]}）" if worst else "未找到图件",
    )


def _check_figure_font_coverage(results: list[dict]) -> None:
    """图件所用的全部汉字必须被所选字体覆盖，避免出现缺字或替代字形。"""
    from matplotlib import font_manager
    from fontTools.ttLib import TTFont

    source = (PROJECT_ROOT / "analysis" / "make_q1_figures.py").read_text(encoding="utf-8")
    characters = sorted({char for char in source if 0x3000 <= ord(char) <= 0x9FFF})
    font_path = font_manager.findfont(
        font_manager.FontProperties(
            family=[
                "Hiragino Sans GB",
                "Songti SC",
                "Heiti TC",
                "Arial Unicode MS",
                "DejaVu Sans",
            ]
        )
    )
    codepoints = set()
    for index in range(4):
        try:
            font = TTFont(font_path, fontNumber=index)
        except Exception:
            break
        cmap = font.getBestCmap() or {}
        codepoints.update(cmap.keys())
        if index == 0 and not font_path.lower().endswith(".ttc"):
            break
    missing = [char for char in characters if ord(char) not in codepoints]
    _record(
        results,
        "图件字体覆盖全部汉字",
        not missing,
        f"字体 {Path(font_path).name}，检查 {len(characters)} 个字符"
        + (f"，缺字 {missing}" if missing else "，无缺字"),
    )


def _check_derived_claims(results: list[dict]) -> None:
    """文档中的派生量必须能由结果文件复算得到（防止未经验证的陈述）。"""
    with (OUTPUTS / "result3_solution.json").open(encoding="utf-8") as file:
        q3 = json.load(file)
    with (OUTPUTS / "result4_solution.json").open(encoding="utf-8") as file:
        q4 = json.load(file)
    t3 = np.array(q3["time_s"], dtype=float)
    c3 = np.array(q3["moisture_dry_basis"], dtype=float)[:, 0]
    t4 = np.array(q4["time_s"], dtype=float)
    c4 = np.array(q4["moisture_dry_basis"], dtype=float)[:, 0]
    grid = np.arange(3600.0, min(t3[-1], t4[-1]), 60.0)
    difference = np.interp(grid, t4, c4) - np.interp(grid, t3, c3)
    index = int(np.nonzero(difference <= 0.0)[0][0])
    t_prev, t_cur = grid[index - 1], grid[index]
    d_prev, d_cur = difference[index - 1], difference[index]
    crossing_hours = (t_prev + (0.0 - d_prev) * (t_cur - t_prev) / (d_cur - d_prev)) / 3600.0
    _record(
        results,
        "派生量：两问圆心曲线交叉时刻",
        abs(crossing_hours - 40.96) <= 0.01,
        f"复算 {crossing_hours:.2f} h",
    )
    text = (PROJECT_ROOT / "docs" / "paper_draft.md").read_text(encoding="utf-8")
    _record(
        results,
        "论文采用复算后的交叉时刻",
        "40.96 h" in text and "15 h" not in text,
        "正文含 40.96 h 且无过期表述" if ("40.96 h" in text and "15 h" not in text) else "正文与复算不一致",
    )
    section_ids = set(
        re.findall(r"^#{2,3}\s*(\d+(?:\.\d+)?)[\.\s]", text, flags=re.MULTILINE)
    )
    references = set(re.findall(r"第\s*(\d+(?:\.\d+)?)\s*节", text))
    dangling = sorted(ref for ref in references if ref not in section_ids)
    _record(
        results,
        "正文交叉引用无悬空",
        not dangling,
        "全部引用有效" if not dangling else f"悬空引用 {dangling}",
    )
    appendix_ids = set(re.findall(r"^##\s*附录\s*([A-Z])", text, flags=re.MULTILINE))
    appendix_refs = set(re.findall(r"附录\s*([A-Z])", text))
    dangling_appendix = sorted(ref for ref in appendix_refs if ref not in appendix_ids)
    _record(
        results,
        "附录引用无悬空",
        not dangling_appendix,
        "全部引用有效" if not dangling_appendix else f"悬空引用 {dangling_appendix}",
    )


def run_checks() -> list[dict]:
    results: list[dict] = []
    for problem in (1, 2, 3, 4):
        _check_workbook(problem, results)
    for problem in (3, 4):
        _check_moisture_semantics(problem, results)
    _check_temperature_bound(results)
    _check_attachments_read_only(results)
    _check_manifests(results)
    _check_headline_numbers(results)
    _check_figures(results)
    _check_documents(results)
    _check_paper_draft(results)
    _check_grayscale_figures(results)
    _check_figure_font_coverage(results)
    _check_derived_claims(results)
    return results


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--strict", action="store_true", help="有 FAIL 时以非零状态退出")
    parser.add_argument("--json", type=Path, default=None, help="把结果写入 JSON")
    args = parser.parse_args()
    results = run_checks()
    print("| 检查项 | 状态 | 详情 |")
    print("| --- | --- | --- |")
    for item in results:
        print(f"| {item['name']} | {item['status']} | {item['detail']} |")
    failures = [item for item in results if item["status"] == "FAIL"]
    print(f"\n合计 {len(results)} 项，PASS {len(results) - len(failures)}，FAIL {len(failures)}")
    if args.json:
        args.json.parent.mkdir(parents=True, exist_ok=True)
        with args.json.open("w", encoding="utf-8") as file:
            json.dump(results, file, ensure_ascii=False, indent=2)
    if args.strict and failures:
        raise SystemExit(1)


if __name__ == "__main__":
    main()

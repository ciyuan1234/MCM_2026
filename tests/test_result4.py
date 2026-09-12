import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from src.config import Q4_OUTPUT_TIME_STEP_S, Q4_SURFACE_HEADER, Q4_TARGET_MOISTURE
from src.export_xlsx import verify_single_sheet_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOLUTION_PATH = PROJECT_ROOT / "outputs" / "result4_solution.json"


def _load_solution():
    with SOLUTION_PATH.open(encoding="utf-8") as file:
        return json.load(file)


def test_result4_workbook_matches_solution() -> None:
    solution = _load_solution()
    output_radius_m = np.arange(0.0, 0.0200001, 0.001)
    radius_m = np.array(solution["radius_m"], dtype=float)
    output_indices = [
        int(np.argmin(np.abs(radius_m - target_radius_m)))
        for target_radius_m in output_radius_m
    ]
    verify_single_sheet_workbook(
        PROJECT_ROOT / "outputs" / "result4.xlsx",
        "Sheet1",
        np.array(solution["time_s"], dtype=float),
        output_radius_m,
        np.array(solution["moisture_dry_basis"], dtype=float)[:, output_indices],
        last_header=Q4_SURFACE_HEADER,
    )


def test_result4_header_uses_material_points_and_moving_surface() -> None:
    workbook = load_workbook(
        PROJECT_ROOT / "outputs" / "result4.xlsx", read_only=True, data_only=True
    )
    try:
        header = tuple(next(workbook["Sheet1"].iter_rows(values_only=True)))
    finally:
        workbook.close()
    assert header[0] == "时间\\到药材中心的距离"
    assert header[-1] == Q4_SURFACE_HEADER
    assert [float(value) for value in header[1:-1]] == [
        round(index * 0.1, 10) for index in range(20)
    ]
    assert len(header) == 22


def test_result4_time_grid_and_termination() -> None:
    solution = _load_solution()
    time_s = np.array(solution["time_s"], dtype=float)
    moisture = np.array(solution["moisture_dry_basis"], dtype=float)
    assert time_s[0] == 0.0
    assert np.allclose(np.diff(time_s), Q4_OUTPUT_TIME_STEP_S, atol=1e-9)
    assert np.isclose(
        float(solution["metadata"]["t_f_grid_s"]), float(time_s[-1]), atol=1e-9
    )
    center = moisture[:, 0]
    assert float(center[-1]) <= Q4_TARGET_MOISTURE
    assert float(center[-2]) > Q4_TARGET_MOISTURE


def test_result4_center_is_spatial_maximum_and_monotone() -> None:
    solution = _load_solution()
    moisture = np.array(solution["moisture_dry_basis"], dtype=float)
    assert np.allclose(moisture.max(axis=1), moisture[:, 0], atol=1e-12)
    assert np.all(np.diff(moisture[:, 0]) <= 1e-12)


def test_result4_metadata_records_shrinkage_and_conservation() -> None:
    solution = _load_solution()
    metadata = solution["metadata"]
    assert abs(float(metadata["radius_min_cm"]) - 1.198) < 1e-9
    assert abs(float(metadata["radius_end_cm"]) - 1.198) < 1e-9
    diagnostics = metadata["diagnostics"]
    assert float(diagnostics["moisture_balance_relative_error"]) < 1e-10
    assert float(diagnostics["linear_system_residual_max"]) < 1e-9

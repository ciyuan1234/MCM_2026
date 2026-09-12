"""Data preparation and read-back verification for spreadsheet outputs."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Sequence

import numpy as np
from openpyxl import load_workbook

from .radial_solver import SolverResult


def write_solution_json(result: SolverResult, path: Path, metadata: dict) -> None:
    payload = {
        "metadata": metadata,
        "time_s": result.time_s.tolist(),
        "radius_m": result.radius_m.tolist(),
        "temperature_c": result.temperature_c.tolist(),
        "moisture_dry_basis": result.moisture_dry_basis.tolist(),
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as file:
        json.dump(payload, file, ensure_ascii=False, indent=2)


def verify_result1_workbook(
    path: Path,
    expected_time_s: np.ndarray,
    expected_radius_m: np.ndarray,
    expected_temperature_c: np.ndarray,
    expected_moisture: np.ndarray,
    tolerance: float = 5e-5,
) -> None:
    verify_two_sheet_workbook(
        path,
        ("温度", "水分浓度"),
        expected_time_s,
        expected_radius_m,
        (expected_temperature_c, expected_moisture),
        tolerance,
    )


def verify_single_sheet_workbook(
    path: Path,
    sheet_name: str,
    expected_time_s: np.ndarray,
    expected_radius_m: np.ndarray,
    expected_values: np.ndarray,
    tolerance: float = 5e-5,
) -> None:
    verify_workbook(
        path,
        (sheet_name,),
        expected_time_s,
        expected_radius_m,
        (expected_values,),
        tolerance,
    )


def verify_two_sheet_workbook(
    path: Path,
    sheet_names: tuple[str, str],
    expected_time_s: np.ndarray,
    expected_radius_m: np.ndarray,
    expected_matrices: tuple[np.ndarray, np.ndarray],
    tolerance: float = 5e-5,
) -> None:
    verify_workbook(
        path,
        sheet_names,
        expected_time_s,
        expected_radius_m,
        expected_matrices,
        tolerance,
    )


def verify_workbook(
    path: Path,
    sheet_names: Sequence[str],
    expected_time_s: np.ndarray,
    expected_radius_m: np.ndarray,
    expected_matrices: Sequence[np.ndarray],
    tolerance: float = 5e-5,
) -> None:
    """Read back a result workbook that holds one or more sheets."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        missing = [name for name in sheet_names if name not in workbook.sheetnames]
        if missing:
            raise AssertionError(
                f"{path.name}: missing sheets {missing}; found {workbook.sheetnames}"
            )
        for sheet_name, expected_values in zip(sheet_names, expected_matrices):
            _verify_sheet(
                workbook[sheet_name],
                expected_time_s,
                expected_radius_m,
                expected_values,
                tolerance,
            )
    finally:
        workbook.close()


def _verify_sheet(
    sheet,
    expected_time_s: np.ndarray,
    expected_radius_m: np.ndarray,
    expected_values: np.ndarray,
    tolerance: float,
) -> None:
    n_rows = expected_time_s.size
    n_columns = expected_radius_m.size
    rows = sheet.iter_rows(
        min_row=1,
        max_row=n_rows + 1,
        max_col=n_columns + 1,
        values_only=True,
    )
    header = tuple(next(rows))
    data_rows = list(rows)

    if len(data_rows) != n_rows:
        raise AssertionError(f"{sheet.title}: expected {n_rows} data rows")
    if header[0] != "时间\\到药材中心的距离":
        raise AssertionError(f"{sheet.title}: unexpected A1 header")

    expected_distances_cm = np.round(expected_radius_m * 100.0, 1)
    actual_distances_cm = np.array([float(value) for value in header[1:]])
    if not np.allclose(actual_distances_cm, expected_distances_cm, atol=1e-12):
        raise AssertionError(f"{sheet.title}: distance headers do not match")

    actual = np.array(data_rows, dtype=float)
    if not np.allclose(actual[:, 0], expected_time_s, atol=tolerance, rtol=0.0):
        raise AssertionError(f"{sheet.title}: time column does not match")
    if not np.allclose(actual[:, 1:], expected_values, atol=tolerance, rtol=0.0):
        maximum_error = float(np.max(np.abs(actual[:, 1:] - expected_values)))
        raise AssertionError(
            f"{sheet.title}: data values do not match; max error={maximum_error:.8g}"
        )

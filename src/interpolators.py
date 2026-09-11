"""Read attachment tables and provide piecewise-linear interpolation."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


@dataclass(frozen=True)
class Attachment1Data:
    time_s: np.ndarray
    air_temperature_c: np.ndarray
    air_moisture_dry_basis: np.ndarray


@dataclass(frozen=True)
class Attachment2Data:
    time_s: np.ndarray
    radius_cm: np.ndarray


class PiecewiseLinear:
    """Piecewise-linear interpolator that holds endpoint values outside the range."""

    def __init__(self, x: np.ndarray, y: np.ndarray) -> None:
        x_arr = np.asarray(x, dtype=float)
        y_arr = np.asarray(y, dtype=float)
        if x_arr.ndim != 1 or y_arr.ndim != 1:
            raise ValueError("interpolation inputs must be one-dimensional")
        if x_arr.size != y_arr.size or x_arr.size < 2:
            raise ValueError("interpolation inputs must have matching lengths >= 2")
        if not np.all(np.isfinite(x_arr)) or not np.all(np.isfinite(y_arr)):
            raise ValueError("interpolation inputs must be finite")
        if np.any(np.diff(x_arr) <= 0.0):
            raise ValueError("interpolation x values must be strictly increasing")
        self.x = x_arr
        self.y = y_arr

    def __call__(self, x_new: np.ndarray | float) -> np.ndarray:
        return np.interp(np.asarray(x_new, dtype=float), self.x, self.y)


def _read_numeric_columns(path: Path, n_columns: int) -> list[tuple[float, ...]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows: list[tuple[float, ...]] = []
        for row in sheet.iter_rows(min_row=2, max_col=n_columns, values_only=True):
            values = tuple(row[:n_columns])
            if any(value is None for value in values):
                raise ValueError(f"{path.name} contains a missing value: {values}")
            try:
                numeric = tuple(float(value) for value in values)
            except (TypeError, ValueError) as exc:
                raise ValueError(f"{path.name} contains a non-numeric value: {values}") from exc
            if not all(np.isfinite(value) for value in numeric):
                raise ValueError(f"{path.name} contains a non-finite value: {values}")
            rows.append(numeric)
    finally:
        workbook.close()

    if len(rows) < 2:
        raise ValueError(f"{path.name} must contain at least two data rows")
    return rows


def _validate_time_column(time_s: np.ndarray, name: str) -> None:
    if np.any(np.diff(time_s) <= 0.0):
        raise ValueError(f"{name} time values must be strictly increasing")
    if np.any(time_s < 0.0):
        raise ValueError(f"{name} time values must be non-negative")


def read_attachment_1(path: Path) -> Attachment1Data:
    rows = _read_numeric_columns(path, 3)
    time_s = np.array([row[0] for row in rows], dtype=float)
    air_temperature_c = np.array([row[1] for row in rows], dtype=float)
    air_moisture = np.array([row[2] for row in rows], dtype=float)
    _validate_time_column(time_s, path.name)
    if np.any(air_moisture < 0.0):
        raise ValueError("attachment 1 contains negative moisture concentration")
    return Attachment1Data(time_s, air_temperature_c, air_moisture)


def read_attachment_2(path: Path) -> Attachment2Data:
    rows = _read_numeric_columns(path, 2)
    time_s = np.array([row[0] for row in rows], dtype=float)
    radius_cm = np.array([row[1] for row in rows], dtype=float)
    _validate_time_column(time_s, path.name)
    if np.any(radius_cm <= 0.0):
        raise ValueError("attachment 2 contains non-positive radius")
    return Attachment2Data(time_s, radius_cm)
